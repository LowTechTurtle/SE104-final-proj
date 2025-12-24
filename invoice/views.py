# Standard library imports
import openpyxl

# Django core imports
from django.urls import reverse
from django.http import HttpResponse # <--- THÊM MỚI
from django.contrib.auth.decorators import login_required # <--- THÊM MỚI

# Authentication and permissions
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.contrib import messages
from django.db import transaction
from django.core.exceptions import ValidationError
from django.http import HttpResponseRedirect

# Class-based views
from django.views.generic import (
    DetailView, CreateView, UpdateView, DeleteView
)

# Third-party packages
from django_tables2 import SingleTableView
from django_tables2.export.views import ExportMixin

# Local app imports
from .models import Invoice
from .tables import InvoiceTable
from .forms import InvoiceForm  
from store.models import Delivery

# ----------------------------------------------------------------------------
# EXISTING CLASS-BASED VIEWS
# ----------------------------------------------------------------------------

class InvoiceListView(LoginRequiredMixin, ExportMixin, SingleTableView):
    """
    View for listing invoices with table export functionality.
    """
    model = Invoice
    table_class = InvoiceTable
    template_name = 'invoice/invoicelist.html'
    context_object_name = 'invoices'
    paginate_by = 10
    table_pagination = False


class InvoiceDetailView(DetailView):
    """
    View for displaying invoice details.
    """
    model = Invoice
    template_name = 'invoice/invoicedetail.html'

    def get_success_url(self):
        return reverse('invoice-detail', kwargs={'slug': self.object.pk})


class InvoiceCreateView(LoginRequiredMixin, CreateView):
    """
    View for creating a new invoice.
    """
    model = Invoice
    template_name = 'invoice/invoicecreate.html'
    form_class = InvoiceForm

    def get_success_url(self):
        return reverse('invoicelist')

    def form_valid(self, form):
        try:
            # Dùng transaction.atomic để đảm bảo: Nếu lỗi thì không lưu gì cả
            with transaction.atomic():
                # 1. Cố gắng lưu Invoice (Lúc này Signal kiểm tra kho sẽ chạy)
                # Nếu không đủ hàng, Signal sẽ "ném" ra ValidationError -> Nhảy xuống except
                self.object = form.save()
                
                # 2. Tạo Delivery tự động (Code cũ của chúng ta)
                Delivery.objects.create(
                    invoice=self.object,
                    location=self.object.customer.address if self.object.customer else "Tại cửa hàng",
                    is_delivered=False
                )

            # 3. Nếu chạy êm đẹp đến đây -> Gửi thông báo Thành công
            messages.success(self.request, "Invoice has been created successfully!")
            return super().form_valid(form)

        except ValidationError as e:
            # 4. Nếu bắt được lỗi (từ Signal) -> Gửi thông báo Lỗi
            # Lấy nội dung lỗi (bỏ cái dấu ngoặc vuông [] đi cho đẹp)
            error_message = e.messages[0] if hasattr(e, 'messages') else str(e)
            messages.error(self.request, error_message)
            
            # Load lại trang hiện tại (không chuyển trang) để hiện lỗi
            return self.render_to_response(self.get_context_data(form=form))


class InvoiceUpdateView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
    """
    View for updating an existing invoice.
    """
    model = Invoice
    template_name = 'invoice/invoiceupdate.html'
    form_class = InvoiceForm
    permission_required = "invoice.change_invoice"

    def get_success_url(self):
        return reverse('invoicelist')

    def test_func(self):
        return self.request.user.is_superuser


class InvoiceDeleteView(LoginRequiredMixin, PermissionRequiredMixin, DeleteView):
    """
    View for deleting an invoice.
    """
    model = Invoice
    template_name = 'invoice/invoicedelete.html'
    success_url = '/products'
    permission_required = "invoice.delete_invoice"

    def get_success_url(self):
        return reverse('invoicelist')

    def test_func(self):
        return self.request.user.is_superuser


# ----------------------------------------------------------------------------
# CUSTOM EXPORT FUNCTION (FIX BUG EXCEL)
# ----------------------------------------------------------------------------

@login_required
def export_invoices(request):
    """
    Export danh sách Invoice ra Excel với đầy đủ thông tin khách hàng, Shipping, Grand Total
    """
    # 1. Tạo file Excel trong bộ nhớ
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Invoices_List.xlsx"'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"

    # 2. Định nghĩa Header (Đầy đủ cột)
    columns = [
        'ID', 
        'Date', 
        'Customer Name', 
        'Contact Number', 
        'Item Name', 
        'Price Per Item', 
        'Quantity', 
        'Total', 
        'Shipping',        # Cột bạn cần
        'Grand Total'      # Cột bạn cần
    ]
    ws.append(columns)

    # Format Header in đậm
    header_font = openpyxl.styles.Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    # 3. Query dữ liệu
    # Dùng select_related để tối ưu hóa truy vấn
    rows = Invoice.objects.all().select_related('customer', 'item').order_by('-date')

    for invoice in rows:
        # -- Xử lý ngày tháng --
        inv_date = invoice.date.strftime('%Y-%m-%d %H:%M') if invoice.date else "-"

        # -- Xử lý Khách hàng (Lấy từ quan hệ Customer) --
        if invoice.customer:
            cust_name = f"{invoice.customer.first_name} {invoice.customer.last_name}"
            cust_phone = invoice.customer.phone if invoice.customer.phone else "-"
        else:
            cust_name = "Guest"
            cust_phone = "-"

        # -- Xử lý Item (Chỉ lấy tên, không lấy thông tin thừa) --
        item_name = invoice.item.name if invoice.item else "Unknown Item"

        # -- Ghi dòng dữ liệu --
        ws.append([
            invoice.id,
            inv_date,
            cust_name,
            cust_phone,
            item_name,
            invoice.price_per_item,
            invoice.quantity,
            invoice.total,
            invoice.shipping,    # Đã bổ sung
            invoice.grand_total  # Đã bổ sung
        ])

    # 4. Auto-fit độ rộng cột
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    wb.save(response)
    return response